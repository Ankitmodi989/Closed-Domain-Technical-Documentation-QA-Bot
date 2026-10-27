"""Evaluate the SQuAD 2.0 QA model on the supplied Excel dev dataset."""

import argparse
import ast
import random
import re
import string
from collections import defaultdict

from openpyxl import load_workbook
from transformers import pipeline


MODEL_NAME = "deepset/roberta-base-squad2"
DEFAULT_DATASET = "data/evaluation/squad_2.0_dev_complete_11873.xlsx"


def normalize(text):
    text = text.lower()
    text = "".join(char for char in text if char not in string.punctuation)
    text = re.sub(r"\b(a|an|the)\b", " ", text)
    return " ".join(text.split())


def exact_match(prediction, references):
    normalized = normalize(prediction)
    return max((normalized == normalize(ref) for ref in references), default=not normalized)


def token_f1(prediction, references):
    pred_tokens = normalize(prediction).split()
    if not references:
        return float(not pred_tokens)
    scores = []
    for reference in references:
        ref_tokens = normalize(reference).split()
        if not pred_tokens or not ref_tokens:
            scores.append(float(pred_tokens == ref_tokens))
            continue
        common = defaultdict(int)
        for token in pred_tokens:
            common[token] += 1
        overlap = 0
        for token in ref_tokens:
            if common[token] > 0:
                overlap += 1
                common[token] -= 1
        if not overlap:
            scores.append(0.0)
        else:
            precision = overlap / len(pred_tokens)
            recall = overlap / len(ref_tokens)
            scores.append(2 * precision * recall / (precision + recall))
    return max(scores)


def read_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        item = dict(zip(headers, values))
        answers = ast.literal_eval(item["answers"]) if item["answers"] else []
        rows.append({
            "context": item["context"],
            "question": item["question"],
            "answers": answers if isinstance(answers, list) else [],
            "impossible": bool(item["is_impossible"]),
        })
    workbook.close()
    return rows


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", default=DEFAULT_DATASET, help="Path to the supplied SQuAD Excel file")
    parser.add_argument("--limit", type=int, default=100, help="Balanced random sample size (default: 100)")
    parser.add_argument("--all", action="store_true", help="Evaluate all rows (slow on CPU)")
    parser.add_argument("--seed", type=int, default=42, help="Random seed for repeatable samples")
    args = parser.parse_args()

    rows = read_rows(args.file)
    if not args.all and args.limit > 0 and args.limit < len(rows):
        randomizer = random.Random(args.seed)
        groups = {True: [], False: []}
        for row in rows:
            groups[row["impossible"]].append(row)
        selected = []
        for impossible, group in groups.items():
            count = round(args.limit * len(group) / len(rows))
            selected.extend(randomizer.sample(group, min(count, len(group))))
        rows = selected

    qa = pipeline("question-answering", model=MODEL_NAME, tokenizer=MODEL_NAME)
    totals = {"all": [0, 0, 0], "answerable": [0, 0, 0], "unanswerable": [0, 0, 0]}
    for index, row in enumerate(rows, start=1):
        prediction = qa(
            question=row["question"],
            context=row["context"],
            handle_impossible_answer=True,
        )["answer"]
        references = [] if row["impossible"] else row["answers"]
        category = "unanswerable" if row["impossible"] else "answerable"
        for key in ("all", category):
            totals[key][0] += exact_match(prediction, references)
            totals[key][1] += token_f1(prediction, references)
            totals[key][2] += 1
        if index % 25 == 0:
            print(f"Processed {index}/{len(rows)} questions")

    print(f"\nModel: {MODEL_NAME}")
    print(f"Examples evaluated: {len(rows)}")
    print("Category       Exact Match     Token F1     Count")
    for category, (em_sum, f1_sum, count) in totals.items():
        em = em_sum / count if count else 0
        f1 = f1_sum / count if count else 0
        print(f"{category:14} {em:11.3f} {f1:13.3f} {count:9}")


if __name__ == "__main__":
    main()
